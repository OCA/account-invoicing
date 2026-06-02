# Copyright 2026 Ecosoft Co., Ltd (https://ecosoft.co.th/)
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl.html)

import base64
from collections import defaultdict
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from odoo import _, fields, models
from odoo.exceptions import UserError


class BillingSummaryWizard(models.TransientModel):
    _name = "billing.summary.wizard"
    _description = "Billing Summary Wizard"

    date_from = fields.Date(required=True)
    date_to = fields.Date(required=True)
    customer_ids = fields.Many2many(
        comodel_name="res.partner",
        string="Customers",
    )
    company_ids = fields.Many2many(
        comodel_name="res.company",
        string="Companies",
        default=lambda self: self.env.company,
        required=True,
    )

    # ── Data helpers ────────────────────────────────────────────────────────────

    def _get_domain(self):
        """Build account.move search domain from wizard fields."""
        domain = [
            ("move_type", "in", ["out_invoice", "out_refund"]),
            ("state", "=", "posted"),
            ("invoice_date", ">=", self.date_from),
            ("invoice_date", "<=", self.date_to),
        ]
        if self.company_ids:
            domain.append(("company_id", "in", self.company_ids.ids))
        if self.customer_ids:
            domain.append(("partner_id", "in", self.customer_ids.ids))
        return domain

    def _get_invoices(self):
        """Fetch posted invoices/credit notes matching the wizard filters."""
        return self.env["account.move"].search(
            self._get_domain(), order="partner_id, invoice_date"
        )

    def _get_invoice_lines(self, invoices):
        """Return all product lines for the given invoices via direct search."""
        return self.env["account.move.line"].search(
            [
                ("move_id", "in", invoices.ids),
                ("display_type", "=", "product"),
                ("product_id", "!=", False),
            ]
        )

    def _get_products(self, lines):
        """Return ordered list of unique products found in the invoice lines."""
        return lines.mapped("product_id").sorted(key=lambda p: p.name or "")

    def _get_report_data(self, lines):
        """
        Build customer × product amount matrix.

        Returns a list of dicts (sorted by partner name):
            [{"partner": res.partner, "amounts": {product_id: float}}, ...]

        Credit notes (out_refund) are subtracted from totals.
        """
        data = {}
        for line in lines:
            partner = line.move_id.partner_id
            if partner.id not in data:
                data[partner.id] = {
                    "partner": partner,
                    "amounts": defaultdict(float),
                }
            sign = -1 if line.move_id.move_type == "out_refund" else 1
            data[partner.id]["amounts"][line.product_id.id] += (
                sign * line.price_subtotal
            )
        return sorted(data.values(), key=lambda r: r["partner"].name or "")

    # ── Style helpers ────────────────────────────────────────────────────────────

    def _get_styles(self):
        """Return a dict of reusable openpyxl style objects."""
        thin = Side(style="thin")
        return {
            "title": Font(bold=True, size=14),
            "label": Font(bold=True, size=11),
            "header": Font(bold=True, size=11),
            "header_fill": PatternFill("solid", fgColor="D9E1F2"),
            "border": Border(left=thin, right=thin, top=thin, bottom=thin),
            "center": Alignment(horizontal="center", vertical="center"),
            "left": Alignment(horizontal="left", vertical="center"),
            "right": Alignment(horizontal="right", vertical="center"),
        }

    # ── Sheet writers ────────────────────────────────────────────────────────────

    def _write_title_section(self, ws, styles):
        """
        Write rows 1–3:
            1A  Billing Summary (Invoice)
            2A  From Date   2B  01-MAR-2026
            3A  To Date     3B  31-MAR-2026
        """
        fmt = "%d-%b-%Y"
        ws["A1"] = "Billing Summary (Invoice)"
        ws["A1"].font = styles["title"]

        ws["A2"] = "From Date"
        ws["A2"].font = styles["label"]
        ws["B2"] = self.date_from.strftime(fmt).upper()

        ws["A3"] = "To Date"
        ws["A3"].font = styles["label"]
        ws["B3"] = self.date_to.strftime(fmt).upper()

    def _write_column_headers(self, ws, products, styles):
        """
        Write row 5 headers:
            A5  Customer Code
            B5  Customer Name
            C5+ one column per product
        """
        headers = ["Customer Code", "Customer Name"] + [p.name or "" for p in products]
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=col, value=text)
            cell.font = styles["header"]
            cell.fill = styles["header_fill"]
            cell.alignment = styles["center"]
            cell.border = styles["border"]

    def _write_data_rows(self, ws, report_data, products, styles):
        """
        Write rows 6+ — one row per customer.
        Show amount per product; use '-' when the customer has no amount.
        """
        for row_idx, row in enumerate(report_data, start=6):
            partner = row["partner"]

            code_cell = ws.cell(row=row_idx, column=1, value=partner.ref or "")
            code_cell.border = styles["border"]
            code_cell.alignment = styles["left"]

            name_cell = ws.cell(row=row_idx, column=2, value=partner.name or "")
            name_cell.border = styles["border"]
            name_cell.alignment = styles["left"]

            for col_idx, product in enumerate(products, start=3):
                amount = row["amounts"].get(product.id, 0.0)
                if amount:
                    cell = ws.cell(row=row_idx, column=col_idx, value=amount)
                    cell.number_format = "#,##0.00"
                    cell.alignment = styles["right"]
                else:
                    cell = ws.cell(row=row_idx, column=col_idx, value="-")
                    cell.alignment = styles["center"]
                cell.border = styles["border"]

    def _auto_fit_columns(self, ws):
        """Adjust each column width to the longest cell value (capped at 40)."""
        for col in ws.columns:
            max_len = max(
                (len(str(cell.value)) for cell in col if cell.value), default=0
            )
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    # ── Main action ──────────────────────────────────────────────────────────────

    def action_export_excel(self):
        """Generate the billing summary Excel file and return a download action."""
        self.ensure_one()

        invoices = self._get_invoices()
        if not invoices:
            raise UserError(_("No posted invoices found for the selected criteria."))

        lines = self._get_invoice_lines(invoices)
        products = self._get_products(lines)
        report_data = self._get_report_data(lines)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Billing Summary"

        styles = self._get_styles()
        self._write_title_section(ws, styles)
        self._write_column_headers(ws, products, styles)
        self._write_data_rows(ws, report_data, products, styles)
        self._auto_fit_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = "billing_summary_{}.xlsx".format(self.date_from.strftime("%Y%m"))
        attachment = self.env["ir.attachment"].create(
            {
                "name": filename,
                "type": "binary",
                "datas": base64.b64encode(output.read()),
                "res_model": self._name,
                "res_id": self.id,
                "mimetype": (
                    "application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"
                ),
            }
        )
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}/{filename}?download=true",
            "target": "new",
        }
